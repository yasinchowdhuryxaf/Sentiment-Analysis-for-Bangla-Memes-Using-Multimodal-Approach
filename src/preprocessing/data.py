"""Read local inputs without modifying original annotations."""
import csv
import json
import os
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
REQUIRED = {"Caption", "Label", "Image_Path"}

def configured_paths():
    dataset = Path(os.environ.get("MEME_DATASET_PATH", REPO_ROOT / "sample_data/synthetic_samples.json"))
    images = Path(os.environ.get("MEME_IMAGE_ROOT", REPO_ROOT / "sample_data"))
    return dataset.expanduser().resolve(), images.expanduser().resolve()

def load_records(path):
    path = Path(path)
    if path.suffix.lower() == ".json":
        records = json.loads(path.read_text(encoding="utf-8"))
    elif path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            records = list(csv.DictReader(handle))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            header = next(rows, ())
            records = [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]
        finally:
            workbook.close()
    else:
        raise ValueError("Expected a JSON, CSV, or XLSX annotation file")
    if not isinstance(records, list):
        raise ValueError("Annotations must contain a list of records")
    for index, record in enumerate(records):
        if not isinstance(record, dict) or not REQUIRED.issubset(record):
            raise ValueError(f"Record {index} must contain {sorted(REQUIRED)}")
        if not isinstance(record["Caption"], str) or not isinstance(record["Image_Path"], str):
            raise ValueError(f"Record {index}: Caption and Image_Path must be text")
        if record["Label"] not in ("Non Offensive", "Offensive"):
            raise ValueError(f"Record {index}: unknown label {record['Label']!r}")
    return records

def resolve_image(root, stored_path):
    root = Path(root).resolve()
    path = (root / stored_path).resolve()
    if not path.is_relative_to(root):
        raise ValueError("Image path must remain inside MEME_IMAGE_ROOT")
    if not path.is_file():
        raise FileNotFoundError(path)
    return path
